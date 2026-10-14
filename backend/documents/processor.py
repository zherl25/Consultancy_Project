"""Document parsing and chunking for PDF, DOCX, XLSX, and plain text files."""

import io
import re
from pathlib import Path
from typing import List

from .exceptions import DocumentProcessingError, UnsupportedFileTypeError


def _split_units(text: str, max_len: int) -> List[str]:
    """Split on sentence ends and line breaks (table rows have no punctuation), then
    hard-wrap anything still longer than max_len at word boundaries."""
    units: List[str] = []
    for piece in re.split(r"(?<=[.!?])\s+|\n+", text):
        piece = piece.strip()
        while len(piece) > max_len:
            cut = piece.rfind(" ", 0, max_len)
            cut = cut if cut > 0 else max_len
            units.append(piece[:cut])
            piece = piece[cut:].strip()
        if piece:
            units.append(piece)
    return units


def _chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[str]:
    text      = re.sub(r"\n{3,}", "\n\n", text.strip())
    sentences = _split_units(text, chunk_size - overlap)

    chunks:  List[str] = []
    current: str       = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 <= chunk_size:
            current = (current + "\n" + sentence).strip()
        else:
            if current:
                chunks.append(current)
            overlap_text = current[-overlap:] if len(current) > overlap else current
            current = (overlap_text + "\n" + sentence).strip()

    if current:
        chunks.append(current)

    return [c for c in chunks if len(c.strip()) > 30]


def _parse_pdf(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF
    doc   = fitz.open(stream=file_bytes, filetype="pdf")
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n\n".join(pages)


def _parse_docx(file_bytes: bytes) -> str:
    """Extract body text in reading order — paragraphs, tables, headers/footers and text boxes."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table

    doc = Document(io.BytesIO(file_bytes))

    def block_lines(container) -> List[str]:
        lines: List[str] = []
        for item in container.iter_inner_content():
            if isinstance(item, Table):
                lines.extend(table_lines(item))
            elif item.text.strip():
                lines.append(item.text.strip())
        return lines

    def table_lines(table: Table) -> List[str]:
        # One line per row, cells joined with " | " so label/value pairs stay together.
        lines: List[str] = []
        for row in table.rows:
            cells, seen = [], set()
            for cell in row.cells:
                if cell._tc in seen:  # merged cells repeat the same underlying element
                    continue
                seen.add(cell._tc)
                text = " ".join(block_lines(cell))
                if text:
                    cells.append(text)
            if cells:
                lines.append(" | ".join(cells))
        return lines

    def unique(lines: List[str], seen: set) -> List[str]:
        out = [ln for ln in lines if ln not in seen]
        seen.update(out)
        return out

    seen: set = set()
    parts = unique(block_lines(doc), seen)

    for section in doc.sections:
        for hf in (section.header, section.footer):
            if not hf.is_linked_to_previous:
                parts += unique(block_lines(hf), seen)

    # Text boxes live inside drawing runs, which paragraph.text does not include.
    for box in doc.element.body.iter(qn("w:txbxContent")):
        text = " ".join(t.text for t in box.iter(qn("w:t")) if t.text).strip()
        if text:
            parts += unique([text], seen)

    return "\n\n".join(parts)


def _parse_xlsx(file_bytes: bytes) -> str:
    from openpyxl import load_workbook
    wb          = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets_text = []
    for sheet_name in wb.sheetnames:
        ws   = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells))
        if rows:
            sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(sheets_text)


def _parse_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


_PARSERS = {
    ".pdf":  _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".txt":  _parse_text,
    ".md":   _parse_text,
}


def process_document(filename: str, file_bytes: bytes) -> List[str]:
    """Parse a document and return text chunks ready for embedding."""
    ext    = Path(filename).suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFileTypeError(f"Unsupported file type: {ext!r}")

    try:
        full_text = parser(file_bytes)
    except (UnsupportedFileTypeError, DocumentProcessingError):
        raise
    except Exception as exc:
        raise DocumentProcessingError(f"Failed to parse '{filename}': {exc}") from exc

    if not full_text.strip():
        raise DocumentProcessingError("Document appears to be empty or could not be parsed.")

    return _chunk_text(full_text)
