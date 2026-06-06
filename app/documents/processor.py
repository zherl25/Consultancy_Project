"""Document parsing and chunking for PDF, DOCX, XLSX, and plain text files."""

import io
import re
from pathlib import Path
from typing import List

from .exceptions import DocumentProcessingError, UnsupportedFileTypeError


def _chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[str]:
    text      = re.sub(r"\n{3,}", "\n\n", text.strip())
    sentences = re.split(r"(?<=[.!?])\s+", text)

    chunks:  List[str] = []
    current: str       = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 <= chunk_size:
            current = (current + " " + sentence).strip()
        else:
            if current:
                chunks.append(current)
            overlap_text = current[-overlap:] if len(current) > overlap else current
            current = (overlap_text + " " + sentence).strip()

    if current:
        chunks.append(current)

    return [c for c in chunks if len(c.strip()) > 30]


def _parse_pdf(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF
    doc   = fitz.open(stream=file_bytes, filetype="pdf")
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n\n".join(pages)


def _parse_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_xlsx(file_bytes: bytes) -> str:
    from openpyxl import load_workbook
    wb          = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets_text = []
    for sheet_name in wb.sheetnames:
        ws   = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells))
        if rows:
            sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(sheets_text)


def _parse_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


_PARSERS = {
    ".pdf":  _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".txt":  _parse_text,
    ".md":   _parse_text,
}


def process_document(filename: str, file_bytes: bytes) -> List[str]:
    """Parse a document and return text chunks ready for embedding."""
    ext    = Path(filename).suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFileTypeError(f"Unsupported file type: {ext!r}")

    try:
        full_text = parser(file_bytes)
    except (UnsupportedFileTypeError, DocumentProcessingError):
        raise
    except Exception as exc:
        raise DocumentProcessingError(f"Failed to parse '{filename}': {exc}") from exc

    if not full_text.strip():
        raise DocumentProcessingError("Document appears to be empty or could not be parsed.")

    return _chunk_text(full_text)
