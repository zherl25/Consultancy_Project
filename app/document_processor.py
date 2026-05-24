"""Document parsing and chunking for PDF, DOCX, XLSX, and plain text files."""

import io
import re
from pathlib import Path
from typing import List


def _chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[str]:
    """Split text into overlapping chunks by sentence boundaries."""
    text = re.sub(r'\n{3,}', '\n\n', text.strip())
    sentences = re.split(r'(?<=[.!?])\s+', text)

    chunks: List[str] = []
    current = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 <= chunk_size:
            current = (current + " " + sentence).strip()
        else:
            if current:
                chunks.append(current)
            # start next chunk with overlap from end of previous
            overlap_text = current[-overlap:] if len(current) > overlap else current
            current = (overlap_text + " " + sentence).strip()

    if current:
        chunks.append(current)

    return [c for c in chunks if len(c.strip()) > 30]


def parse_pdf(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = []
    for page in doc:
        pages.append(page.get_text())
    doc.close()
    return "\n\n".join(pages)


def parse_docx(file_bytes: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def parse_xlsx(file_bytes: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheets_text = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells))
        if rows:
            sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(sheets_text)


def parse_text(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".txt": parse_text,
    ".md": parse_text,
}


def process_document(filename: str, file_bytes: bytes) -> List[str]:
    """Parse a document and return a list of text chunks."""
    ext = Path(filename).suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext}")

    full_text = parser(file_bytes)
    if not full_text.strip():
        raise ValueError("Document appears to be empty or could not be parsed.")

    return _chunk_text(full_text)
